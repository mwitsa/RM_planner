"""Shared low-level Excel reading helpers for inventory extractors.

Supports legacy .xls (via xlrd), .xlsx/.xlsm (via openpyxl), and the binary
.xlsb format (via pyxlsb), since reports arrive in whichever of these the
source system happened to export.
"""

from __future__ import annotations

import json
from dataclasses import asdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterable, TypeVar

# Excel's day-0 epoch (accounts for its 1900 leap-year bug); pyxlsb returns
# raw serial numbers for date cells instead of resolving them like openpyxl.
_EXCEL_EPOCH = date(1899, 12, 30)

DEFAULT_SUPPORTED_EXTENSIONS = {".xls", ".xlsx", ".xlsm", ".xlsb"}

_RecordT = TypeVar("_RecordT")


def validated_path(
    workbook_path: str | Path,
    supported_extensions: set[str] = DEFAULT_SUPPORTED_EXTENSIONS,
) -> Path:
    path = Path(workbook_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")
    if path.suffix.lower() not in supported_extensions:
        extensions = ", ".join(sorted(supported_extensions))
        raise ValueError(f"Please select a workbook with one of these extensions: {extensions}.")
    return path


def read_sheet_names(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            return list(workbook.sheet_names())
        finally:
            workbook.release_resources()

    if suffix == ".xlsb":
        import pyxlsb

        with pyxlsb.open_workbook(str(path)) as workbook:
            return list(workbook.sheets)

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def read_sheet_rows(path: Path, sheet_name: str) -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(str(path))
        if sheet_name not in workbook.sheet_names():
            raise ValueError(f"Worksheet not found: {sheet_name}")
        sheet = workbook.sheet_by_name(sheet_name)
        return [
            [sheet.cell_value(row, column) for column in range(sheet.ncols)]
            for row in range(sheet.nrows)
        ]

    if suffix == ".xlsb":
        import pyxlsb

        with pyxlsb.open_workbook(str(path)) as workbook:
            if sheet_name not in workbook.sheets:
                raise ValueError(f"Worksheet not found: {sheet_name}")
            with workbook.get_sheet(sheet_name) as sheet:
                return [[cell.v for cell in row] for row in sheet.rows()]

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet_name}")
        worksheet = workbook[sheet_name]
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def choose_default_sheet(sheets: list[str], keyword: str) -> str:
    """Return the detail sheet for `keyword` (e.g. "package"), not its pivot summary.

    Prefers a sheet matching both "data" and `keyword` (e.g. "Data package")
    over one that only matches `keyword` (e.g. a "sum Package" pivot table).
    """

    keyword_fold = keyword.casefold()
    for name in sheets:
        folded = name.casefold()
        if "data" in folded and keyword_fold in folded:
            return name
    for name in sheets:
        if keyword_fold in name.casefold():
            return name
    return sheets[0] if sheets else ""


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value).strip()


def cell_float(value: object) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def cell_date_text(value: object) -> str:
    """Format a date cell as DD/MM/YYYY, handling either a real date/datetime
    (openpyxl resolves formatted cells to these) or a raw Excel serial number
    (pyxlsb does not resolve cell formatting, so dates arrive as floats)."""

    if value is None or value == "":
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%d/%m/%Y")
    try:
        serial = float(value)
    except (TypeError, ValueError):
        return cell_text(value)
    if serial <= 0:
        return cell_text(value)
    return (_EXCEL_EPOCH + timedelta(days=int(serial))).strftime("%d/%m/%Y")


def save_records(
    store_path: str | Path,
    records: Iterable[object],
    *,
    version: int,
    source_file: str = "",
    source_sheet: str = "",
) -> None:
    """Persist extracted records (any flat dataclass) so they survive a restart."""

    path = Path(store_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": version,
        "source_file": source_file,
        "source_sheet": source_sheet,
        "records": [asdict(record) for record in records],
    }
    temporary_path = path.with_suffix(f"{path.suffix}.tmp")
    try:
        with temporary_path.open("w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        temporary_path.replace(path)
    except OSError as exc:
        raise ValueError(f"Could not save data: {exc}") from exc


def load_records(
    store_path: str | Path, record_type: type[_RecordT],
) -> tuple[list[_RecordT], str, str]:
    """Return (records, source_file, source_sheet); ([], "", "") if never saved."""

    path = Path(store_path)
    if not path.exists():
        return [], "", ""
    try:
        with path.open(encoding="utf-8") as handle:
            payload = json.load(handle)
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Could not read saved data: {exc}") from exc
    if not isinstance(payload, dict) or not isinstance(payload.get("records"), list):
        raise ValueError("Saved data has an invalid structure.")
    records = [record_type(**row) for row in payload["records"]]
    return records, payload.get("source_file", ""), payload.get("source_sheet", "")
