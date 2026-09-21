"""Reader and validation model for the shrimp-size assortment master."""

from __future__ import annotations

from dataclasses import dataclass
import math
from pathlib import Path
from typing import Any

@dataclass(frozen=True, slots=True)
class AssortmentTable:
    """Yield distribution from each base shrimp size to actual output sizes."""

    source_path: Path
    sheet_name: str
    base_sizes: tuple[str, ...]
    output_sizes: tuple[str, ...]
    percentages: tuple[tuple[float, ...], ...]
    column_totals: tuple[float, ...]

    @property
    def invalid_base_sizes(self) -> tuple[tuple[str, float], ...]:
        """Return base sizes whose output distribution does not total 100%."""

        return tuple(
            (base_size, total)
            for base_size, total in zip(self.base_sizes, self.column_totals)
            if abs(total - 1.0) > 0.0001
        )


@dataclass(frozen=True, slots=True)
class AssortmentPredictionEntry:
    """Predicted output weight for one actual shrimp-size range."""

    output_size: str
    percentage: float
    weight: int


def predict_assortment(
    table: AssortmentTable,
    harvest_size: str,
    total_weight: float,
) -> tuple[AssortmentPredictionEntry, ...]:
    """Split a harvest weight using one base-size distribution from the master."""

    requested_size = str(harvest_size).strip()
    if not requested_size:
        raise ValueError("Enter a harvest size.")

    lookup = {base_size.casefold(): index for index, base_size in enumerate(table.base_sizes)}
    candidates = [requested_size]
    if not requested_size.casefold().startswith("s."):
        candidates.append(f"S.{requested_size}")
    base_index = next(
        (lookup[candidate.casefold()] for candidate in candidates if candidate.casefold() in lookup),
        None,
    )
    if base_index is None:
        available = ", ".join(table.base_sizes)
        raise ValueError(
            f"Harvest size {requested_size} was not found in the assortment master. "
            f"Available sizes: {available}"
        )

    try:
        weight = float(total_weight)
    except (TypeError, ValueError) as exc:
        raise ValueError("Weight (kg) must be a number.") from exc
    if not math.isfinite(weight) or weight <= 0:
        raise ValueError("Weight (kg) must be greater than zero.")

    distribution_total = table.column_totals[base_index]
    if abs(distribution_total - 1.0) > 0.0001:
        raise ValueError(
            f"The {table.base_sizes[base_index]} distribution totals "
            f"{distribution_total:.2%}, not 100%. Correct the master data before predicting."
        )

    target_weight = math.floor(weight + 0.5)
    if target_weight < 1:
        raise ValueError("Weight (kg) must round to at least 1 kg.")

    distribution = [
        (
            output_size,
            table.percentages[row_index][base_index],
        )
        for row_index, output_size in enumerate(table.output_sizes)
        if table.percentages[row_index][base_index] > 0
    ]
    if not distribution:
        raise ValueError(f"The {table.base_sizes[base_index]} distribution has no output percentages.")

    # Allocate whole kilograms with the largest-remainder method. This keeps
    # every generated row integer-valued while preserving the rounded harvest
    # total exactly.
    raw_weights = [
        target_weight * percentage / distribution_total
        for _output_size, percentage in distribution
    ]
    whole_weights = [math.floor(raw_weight) for raw_weight in raw_weights]
    kilograms_left = target_weight - sum(whole_weights)
    remainder_order = sorted(
        range(len(raw_weights)),
        key=lambda index: (raw_weights[index] - whole_weights[index], -index),
        reverse=True,
    )
    for index in remainder_order[:kilograms_left]:
        whole_weights[index] += 1

    return tuple(
        AssortmentPredictionEntry(
            output_size=output_size,
            percentage=percentage,
            weight=whole_weight,
        )
        for (output_size, percentage), whole_weight in zip(distribution, whole_weights)
        if whole_weight > 0
    )


def load_assortment(workbook_path: str | Path) -> AssortmentTable:
    """Load the first worksheet as an assortment percentage matrix.

    Row 1 contains the base harvested shrimp sizes from column B onward.
    Column A contains actual output-size ranges. Matrix values are stored as
    decimal percentages (for example, 0.404 means 40.4%).
    """

    # Avoid paying openpyxl's import cost until the user needs the assortment
    # workbook.  This keeps the initial application launch responsive.
    from openpyxl import load_workbook

    path = Path(workbook_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Assortment master file not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Assortment master must be an .xlsx or .xlsm file.")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if not workbook.sheetnames:
            raise ValueError("Assortment workbook does not contain a worksheet.")
        worksheet = workbook[workbook.sheetnames[0]]

        base_sizes: list[str] = []
        for column in range(2, worksheet.max_column + 1):
            value = worksheet.cell(1, column).value
            if _is_blank(value):
                break
            base_sizes.append(str(value).strip())
        if not base_sizes:
            raise ValueError("No base shrimp sizes were found in row 1 from column B onward.")
        if len(set(base_sizes)) != len(base_sizes):
            raise ValueError("Base shrimp-size headers must be unique.")

        output_sizes: list[str] = []
        percentage_rows: list[tuple[float, ...]] = []
        for row_number in range(2, worksheet.max_row + 1):
            output_value = worksheet.cell(row_number, 1).value
            if _is_blank(output_value):
                continue
            output_size = str(output_value).strip()
            if output_size in output_sizes:
                raise ValueError(f"Duplicate actual output size: {output_size}")

            percentages = tuple(
                _parse_percentage(worksheet.cell(row_number, column).value, row_number, column)
                for column in range(2, len(base_sizes) + 2)
            )
            output_sizes.append(output_size)
            percentage_rows.append(percentages)

        if not output_sizes:
            raise ValueError("No actual output-size rows were found in column A.")

        column_totals = tuple(
            sum(row[column_index] for row in percentage_rows)
            for column_index in range(len(base_sizes))
        )
        return AssortmentTable(
            source_path=path,
            sheet_name=worksheet.title,
            base_sizes=tuple(base_sizes),
            output_sizes=tuple(output_sizes),
            percentages=tuple(percentage_rows),
            column_totals=column_totals,
        )
    finally:
        workbook.close()


def _parse_percentage(value: Any, row: int, column: int) -> float:
    if _is_blank(value):
        return 0.0
    if isinstance(value, bool):
        raise ValueError(f"Invalid percentage at row {row}, column {column}.")
    if isinstance(value, (int, float)):
        percentage = float(value)
    elif isinstance(value, str):
        cleaned = value.strip()
        try:
            if cleaned.endswith("%"):
                percentage = float(cleaned[:-1].strip()) / 100
            else:
                percentage = float(cleaned)
        except ValueError as exc:
            raise ValueError(f"Invalid percentage at row {row}, column {column}: {value}") from exc
    else:
        raise ValueError(f"Invalid percentage at row {row}, column {column}: {value}")

    if not 0 <= percentage <= 1:
        raise ValueError(
            f"Percentage at row {row}, column {column} must be between 0% and 100%."
        )
    return percentage


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
