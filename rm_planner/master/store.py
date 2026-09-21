"""Extraction and local persistence for the packaging/ingredient master data."""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

STORE_VERSION = 1
MATERIAL_COLUMN = 1
DESCRIPTION_COLUMN = 2
COMPONENT_NAME_COLUMN = 3
MRP_CONTROLLER_COLUMN = 4
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}


@dataclass(frozen=True, slots=True)
class MasterComponentRecord:
    """One Material -> packaging/ingredient component line from the master file."""

    material: str
    description: str
    component_name: str
    mrp_controller: str


def list_sheets(workbook_path: str | Path) -> list[str]:
    """Return workbook sheet names without modifying the workbook."""

    from openpyxl import load_workbook

    path = _validated_path(workbook_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def extract_master_data(
    workbook_path: str | Path,
    sheet_name: str | None = None,
) -> list[MasterComponentRecord]:
    """Read every non-blank row below the header as a component record.

    Column layout: A = Material, B = Description, C = Component name,
    D = Component MRP Controller.
    """

    from openpyxl import load_workbook

    path = _validated_path(workbook_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        selected_sheet = sheet_name or workbook.sheetnames[0]
        if selected_sheet not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {selected_sheet}")

        worksheet = workbook[selected_sheet]
        records: list[MasterComponentRecord] = []
        for row in worksheet.iter_rows(
            min_row=2,
            max_col=MRP_CONTROLLER_COLUMN,
            values_only=True,
        ):
            if all(_is_blank(value) for value in row):
                continue
            material, description, component_name, mrp_controller = _padded(row, 4)
            records.append(
                MasterComponentRecord(
                    material=_clean_text(material),
                    description=_clean_text(description),
                    component_name=_clean_text(component_name),
                    mrp_controller=_clean_text(mrp_controller),
                )
            )
        return records
    finally:
        workbook.close()


def save_master_data(
    store_path: str | Path,
    records: Iterable[MasterComponentRecord],
) -> Path:
    path = Path(store_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": STORE_VERSION,
        "records": [
            {
                "material": record.material,
                "description": record.description,
                "component_name": record.component_name,
                "mrp_controller": record.mrp_controller,
            }
            for record in records
        ],
    }
    temporary_path = path.with_suffix(f"{path.suffix}.tmp")
    try:
        with temporary_path.open("w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
            handle.write("\n")
        temporary_path.replace(path)
    except OSError as exc:
        raise ValueError(f"Could not save master data: {exc}") from exc
    return path


def load_master_data(store_path: str | Path) -> list[MasterComponentRecord]:
    path = Path(store_path)
    if not path.exists():
        return []
    try:
        with path.open(encoding="utf-8") as handle:
            payload = json.load(handle)
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Could not read saved master data: {exc}") from exc
    if not isinstance(payload, dict) or not isinstance(payload.get("records"), list):
        raise ValueError("Saved master data file has an invalid structure.")

    records: list[MasterComponentRecord] = []
    for raw in payload["records"]:
        if not isinstance(raw, dict):
            raise ValueError("Saved master data file contains an invalid record.")
        records.append(
            MasterComponentRecord(
                material=str(raw.get("material", "")),
                description=str(raw.get("description", "")),
                component_name=str(raw.get("component_name", "")),
                mrp_controller=str(raw.get("mrp_controller", "")),
            )
        )
    return records


def _validated_path(workbook_path: str | Path) -> Path:
    path = Path(workbook_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("Please select an .xlsx or .xlsm workbook.")
    return path


def _padded(row: tuple[Any, ...], length: int) -> tuple[Any, ...]:
    if len(row) >= length:
        return row[:length]
    return row + (None,) * (length - len(row))


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
