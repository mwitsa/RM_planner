"""Core extraction logic for CPF production-plan workbooks.

The workbook is treated as an input only.  This module deliberately has no UI
dependencies so that future planning rules can reuse the same structured data.
"""

from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

PRODUCTION_DATE_COLUMN = 3  # C
COUNTRY_COLUMN = 8  # H
CUSTOMER_COLUMN = 10  # J
GROUP_1_COLUMN = 11  # K
GROUP_2_COLUMN = 12  # L
PACKAGING_COLUMN = 15  # O
WONTONS_PER_CUP_COLUMN = 17  # Q
HO_WEIGHT_FACTOR_COLUMN = 18  # R
RM_SIZE_COLUMN = 19  # S
SOUP_COLUMN = 21  # U
ORDER_UNIT_COLUMN = 22  # V
ORDER_CUPS_COLUMN = 35  # AI
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}
ORDER_EXPORT_FIELDS = (
    "date",
    "month",
    "year",
    "country",
    "customer_name",
    "group_1",
    "group_2",
    "packaging",
    "rm_size",
    "soup",
    "wontons_per_cup",
    "order_unit",
    "order_cups",
    "cups_per_unit",
    "total_wontons",
    "ho_weight_kg",
    "production",
)


@dataclass(slots=True)
class OrderRecord:
    """One normalized order extracted from the production plan."""

    date: str
    month: str
    year: str
    country: str
    customer_name: str
    group_1: str
    group_2: str
    packaging: str
    soup: str
    order_unit: int | float
    order_cups: int | float | None
    cups_per_unit: float | None
    rm_size: str = ""
    wontons_per_cup: int | float | None = None
    ho_weight_kg: int | float | None = None
    production: int | float | None = None
    record_id: str = ""
    order_no: str = ""

    @property
    def month_key(self) -> str:
        """Return YYYY-MM for filtering while exports keep separate columns."""

        return f"{self.year}-{self.month}"

    @property
    def total_wontons(self) -> int | float | None:
        """Return Order (ถ้วย) multiplied by ลูกเกี๊ยว/ถ้วย when both exist."""

        return _multiply_optional(self.order_cups, self.wontons_per_cup)


@dataclass(frozen=True, slots=True)
class ExtractionIssue:
    """A non-empty source row that could not become a complete order."""

    source_row: int
    reason: str
    production_date: str
    customer_name: str
    order_unit: str


@dataclass(frozen=True, slots=True)
class ExtractionResult:
    records: list[OrderRecord]
    issues: list[ExtractionIssue]
    sheet_name: str
    header_row: int


def list_sheets(workbook_path: str | Path) -> list[str]:
    """Return workbook sheet names without modifying the workbook."""

    path = _validated_path(workbook_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def choose_default_sheet(workbook_path: str | Path) -> str:
    """Select the sheet whose C/J/V headers most closely match an order table."""

    path = _validated_path(workbook_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        scored: list[tuple[int, int, str]] = []
        for position, worksheet in enumerate(workbook.worksheets):
            header_row, score = _find_header_row(worksheet)
            visible_bonus = 1 if worksheet.sheet_state == "visible" else 0
            scored.append((score * 10 + visible_bonus, -position, worksheet.title))
        if not scored:
            raise ValueError("The workbook does not contain any worksheets.")
        return max(scored)[2]
    finally:
        workbook.close()


def extract_orders(
    workbook_path: str | Path,
    sheet_name: str | None = None,
) -> ExtractionResult:
    """Extract complete C/J/V order rows into a stable, normalized schema.

    Column C becomes separate zero-padded date, month, and year fields. Column J
    becomes the customer name. Column V becomes the order quantity in units,
    column Q becomes ลูกเกี๊ยว/ถ้วย, and column AI becomes the order quantity
    in cups. Cups per unit is derived by dividing column AI by column V, while
    จำนวนเกี๊ยว is derived by multiplying column AI by column Q. น้ำหนัก HO
    (kg) is จำนวนเกี๊ยว multiplied by column R, then divided by 0.54.
    Rows with no selected values are ignored; other incomplete rows are reported
    as issues rather than silently converted to orders.
    """

    path = _validated_path(workbook_path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        selected_sheet = sheet_name or _choose_default_sheet_from_workbook(workbook)
        if selected_sheet not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {selected_sheet}")

        worksheet = workbook[selected_sheet]
        header_row, _ = _find_header_row(worksheet)
        records: list[OrderRecord] = []
        issues: list[ExtractionIssue] = []

        for source_row, row in enumerate(
            worksheet.iter_rows(
                min_row=header_row + 1,
                max_col=ORDER_CUPS_COLUMN,
                values_only=True,
            ),
            start=header_row + 1,
        ):
            raw_date = _cell_value(row, PRODUCTION_DATE_COLUMN)
            raw_country = _cell_value(row, COUNTRY_COLUMN)
            raw_customer = _cell_value(row, CUSTOMER_COLUMN)
            raw_group_1 = _cell_value(row, GROUP_1_COLUMN)
            raw_group_2 = _cell_value(row, GROUP_2_COLUMN)
            raw_packaging = _cell_value(row, PACKAGING_COLUMN)
            raw_wontons_per_cup = _cell_value(row, WONTONS_PER_CUP_COLUMN)
            raw_ho_weight_factor = _cell_value(row, HO_WEIGHT_FACTOR_COLUMN)
            raw_rm_size = _cell_value(row, RM_SIZE_COLUMN)
            raw_soup = _cell_value(row, SOUP_COLUMN)
            raw_unit = _cell_value(row, ORDER_UNIT_COLUMN)
            raw_cups = _cell_value(row, ORDER_CUPS_COLUMN)

            if all(_is_blank(value) for value in (raw_date, raw_customer, raw_unit)):
                continue

            production_period = _parse_production_period(raw_date)
            customer = _clean_text(raw_customer)
            order_unit = _parse_number(raw_unit)
            ho_weight_factor = _parse_number(raw_ho_weight_factor)
            order_cups = _parse_number(raw_cups)
            wontons_per_cup = _parse_number(raw_wontons_per_cup)

            missing: list[str] = []
            if production_period is None:
                missing.append("invalid or missing production date/month (column C)")
            if not customer:
                missing.append("missing customer (column J)")
            if order_unit is None:
                missing.append("invalid or missing order quantity in units (column V)")
            if not _is_blank(raw_cups) and order_cups is None:
                missing.append("invalid order quantity in cups (column AI)")
            if not _is_blank(raw_wontons_per_cup) and wontons_per_cup is None:
                missing.append("invalid ลูกเกี๊ยว/ถ้วย (column Q)")
            if not _is_blank(raw_ho_weight_factor) and ho_weight_factor is None:
                missing.append("invalid น้ำหนัก HO factor (column R)")

            if missing:
                issues.append(
                    ExtractionIssue(
                        source_row=source_row,
                        reason="; ".join(missing),
                        production_date=_display_value(raw_date),
                        customer_name=_display_value(raw_customer),
                        order_unit=_display_value(raw_unit),
                    )
                )
                continue

            records.append(
                OrderRecord(
                    date=production_period[0],
                    month=production_period[1],
                    year=production_period[2],
                    country=_clean_text(raw_country),
                    customer_name=customer,
                    group_1=_clean_text(raw_group_1),
                    group_2=_clean_text(raw_group_2),
                    packaging=_clean_text(raw_packaging),
                    rm_size=_clean_text(raw_rm_size),
                    soup=_clean_text(raw_soup),
                    wontons_per_cup=wontons_per_cup,
                    order_unit=order_unit,
                    order_cups=order_cups,
                    cups_per_unit=_divide_optional(order_cups, order_unit),
                    ho_weight_kg=_calculate_ho_weight(
                        order_cups,
                        wontons_per_cup,
                        ho_weight_factor,
                    ),
                    record_id=f"{selected_sheet}:{source_row}",
                )
            )

        return ExtractionResult(
            records=records,
            issues=issues,
            sheet_name=selected_sheet,
            header_row=header_row,
        )
    finally:
        workbook.close()


def export_csv(records: Iterable[OrderRecord], output_path: str | Path) -> Path:
    """Export records as Excel-friendly UTF-8 CSV."""

    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=ORDER_EXPORT_FIELDS)
        writer.writeheader()
        writer.writerows(_record_payload(record) for record in records)
    return destination


def export_json(records: Iterable[OrderRecord], output_path: str | Path) -> Path:
    """Export records as a UTF-8 JSON array with Thai text preserved."""

    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    payload = [_record_payload(record) for record in records]
    with destination.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)
    return destination


def _record_payload(record: OrderRecord) -> dict[str, Any]:
    return {field: getattr(record, field) for field in ORDER_EXPORT_FIELDS}


def _validated_path(workbook_path: str | Path) -> Path:
    path = Path(workbook_path).expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Workbook not found: {path}")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("Please select an .xlsx or .xlsm workbook.")
    return path


def _choose_default_sheet_from_workbook(workbook: Any) -> str:
    scored: list[tuple[int, int, str]] = []
    for position, worksheet in enumerate(workbook.worksheets):
        _, score = _find_header_row(worksheet)
        visible_bonus = 1 if worksheet.sheet_state == "visible" else 0
        scored.append((score * 10 + visible_bonus, -position, worksheet.title))
    if not scored:
        raise ValueError("The workbook does not contain any worksheets.")
    return max(scored)[2]


def _find_header_row(worksheet: Any, scan_rows: int = 20) -> tuple[int, int]:
    best_row = 4
    best_score = 0
    for row_number in range(1, min(worksheet.max_row, scan_rows) + 1):
        date_header = _normalized_header(worksheet.cell(row_number, PRODUCTION_DATE_COLUMN).value)
        customer_header = _normalized_header(worksheet.cell(row_number, CUSTOMER_COLUMN).value)
        unit_header = _normalized_header(worksheet.cell(row_number, ORDER_UNIT_COLUMN).value)
        score = 0
        if any(token in date_header for token in ("date", "เดือน", "แผนผลิต")):
            score += 1
        if any(token in customer_header for token in ("customer", "ลูกค้า")):
            score += 1
        if any(token in unit_header for token in ("qty", "volume", "จำนวน")):
            score += 1
        if score > best_score:
            best_row, best_score = row_number, score
    return best_row, best_score


def _cell_value(row: tuple[Any, ...], one_based_column: int) -> Any:
    index = one_based_column - 1
    return row[index] if index < len(row) else None


def _normalized_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _parse_production_period(value: Any) -> tuple[str, str, str] | None:
    """Return (day, month, year); day is blank for month-only orders."""

    if isinstance(value, datetime):
        parsed = value.date()
        return f"{parsed.day:02d}", f"{parsed.month:02d}", f"{parsed.year:04d}"
    if isinstance(value, date):
        return f"{value.day:02d}", f"{value.month:02d}", f"{value.year:04d}"
    if isinstance(value, str):
        cleaned = value.strip().lstrip("'").strip()
        month_only = re.fullmatch(r"(0?[1-9]|1[0-2])[-/](\d{4})", cleaned)
        if month_only:
            month = int(month_only.group(1))
            year = int(month_only.group(2))
            if year > 2400:
                year -= 543
            return "", f"{month:02d}", f"{year:04d}"
        for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
            try:
                parsed = datetime.strptime(cleaned, date_format).date()
                if parsed.year > 2400:  # Thai Buddhist Era year
                    parsed = parsed.replace(year=parsed.year - 543)
                return f"{parsed.day:02d}", f"{parsed.month:02d}", f"{parsed.year:04d}"
            except ValueError:
                continue
    return None


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _parse_number(value: Any) -> int | float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        numeric = float(value)
    elif isinstance(value, str):
        cleaned = value.strip().replace(",", "")
        if not cleaned:
            return None
        try:
            numeric = float(cleaned)
        except ValueError:
            return None
    else:
        return None
    if numeric != numeric or numeric in (float("inf"), float("-inf")):
        return None
    return int(numeric) if numeric.is_integer() else numeric


def _divide_optional(
    numerator: int | float | None,
    denominator: int | float,
) -> float | None:
    if numerator is None or denominator == 0:
        return None
    return numerator / denominator


def _multiply_optional(
    first: int | float | None,
    second: int | float | None,
) -> int | float | None:
    if first is None or second is None:
        return None
    result = first * second
    return int(result) if float(result).is_integer() else result


def _calculate_ho_weight(
    order_cups: int | float | None,
    wontons_per_cup: int | float | None,
    ho_weight_factor: int | float | None,
) -> int | float | None:
    total_wontons = _multiply_optional(order_cups, wontons_per_cup)
    weighted_total = _multiply_optional(total_wontons, ho_weight_factor)
    if weighted_total is None:
        return None
    result = weighted_total / 0.54
    return int(result) if float(result).is_integer() else result


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value)
