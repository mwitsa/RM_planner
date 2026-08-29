from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from assortment import load_assortment, predict_assortment


class AssortmentTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary_directory.cleanup)
        self.workbook_path = Path(self.temporary_directory.name) / "assortment.xlsx"

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Assortment"
        sheet.append(["Actual output", "S.43", "S.66"])
        sheet.append(["31-35", 0.4, "20%"])
        sheet.append(["36-40", 0.6, 0.8])
        workbook.save(self.workbook_path)

    def test_loads_assortment_matrix_and_totals(self) -> None:
        table = load_assortment(self.workbook_path)

        self.assertEqual(table.sheet_name, "Assortment")
        self.assertEqual(table.base_sizes, ("S.43", "S.66"))
        self.assertEqual(table.output_sizes, ("31-35", "36-40"))
        self.assertEqual(table.percentages, ((0.4, 0.2), (0.6, 0.8)))
        self.assertEqual(table.column_totals, (1.0, 1.0))
        self.assertEqual(table.invalid_base_sizes, ())

    def test_reports_distribution_that_does_not_total_100_percent(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Actual output", "S.43"])
        sheet.append(["31-35", 0.4])
        workbook.save(self.workbook_path)

        table = load_assortment(self.workbook_path)
        self.assertEqual(table.invalid_base_sizes, (("S.43", 0.4),))

    def test_rejects_percentage_over_100_percent(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Actual output", "S.43"])
        sheet.append(["31-35", 1.1])
        workbook.save(self.workbook_path)

        with self.assertRaises(ValueError):
            load_assortment(self.workbook_path)

    def test_predicts_output_weights_from_plain_harvest_size(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Actual output", "S.74"])
        sheet.append(["61-65", 0.25])
        sheet.append(["66-70", 0.75])
        workbook.save(self.workbook_path)

        predictions = predict_assortment(load_assortment(self.workbook_path), "74", 1_000)

        self.assertEqual(
            [(item.output_size, item.percentage, item.weight) for item in predictions],
            [("61-65", 0.25, 250), ("66-70", 0.75, 750)],
        )
        self.assertTrue(all(isinstance(item.weight, int) for item in predictions))
        self.assertEqual(sum(item.weight for item in predictions), 1_000)

    def test_prediction_accepts_full_master_header(self) -> None:
        predictions = predict_assortment(load_assortment(self.workbook_path), "s.43", 100)
        self.assertEqual([item.weight for item in predictions], [40, 60])

    def test_prediction_allocates_whole_kg_and_reconciles_to_input_total(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Actual output", "S.74"])
        sheet.append(["61-65", 0.333333333])
        sheet.append(["66-70", 0.333333333])
        sheet.append(["71-75", 0.333333334])
        workbook.save(self.workbook_path)

        predictions = predict_assortment(load_assortment(self.workbook_path), "74", 1_000)

        self.assertEqual([item.weight for item in predictions], [333, 333, 334])
        self.assertEqual(sum(item.weight for item in predictions), 1_000)

    def test_prediction_rounds_fractional_total_weight_to_nearest_kg(self) -> None:
        predictions = predict_assortment(load_assortment(self.workbook_path), "43", 100.6)
        self.assertEqual(sum(item.weight for item in predictions), 101)

    def test_prediction_rejects_unknown_size_and_invalid_weight(self) -> None:
        table = load_assortment(self.workbook_path)
        with self.assertRaisesRegex(ValueError, "not found"):
            predict_assortment(table, "74", 100)
        with self.assertRaisesRegex(ValueError, "greater than zero"):
            predict_assortment(table, "43", 0)

    def test_prediction_rejects_incomplete_distribution(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Actual output", "S.74"])
        sheet.append(["61-65", 0.75])
        workbook.save(self.workbook_path)

        with self.assertRaisesRegex(ValueError, "not 100%"):
            predict_assortment(load_assortment(self.workbook_path), "74", 100)


if __name__ == "__main__":
    unittest.main()
